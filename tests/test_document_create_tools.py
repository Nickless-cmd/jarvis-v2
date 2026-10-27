"""Dokument-oprettelse: findes, kan nås, og laver en fil der kan ÅBNES.

## Målingen der udløste filen (27/10-2026)

Bjørn bad om en sammenligning med ChatGPT's evner i chat mode. Målt:
huset kunne LÆSE fem Google-formater gennem connectorer, men der fandtes
**nul** værktøjer der kunne oprette en fil. `pdf_read` fandtes;
`pdf_create` gjorde ikke. Tre af kategorierne — regneark, dokumenter,
præsentationer — var derfor halve: man kunne se indhold, ikke lave det.

## Hvad testen værner

Tre ting, og de er ikke den samme:

1. **Findes** — værktøjet står i kataloget. En handler uden schema er usynlig.
2. **Kan nås** — det slipper gennem scope-porten for den rolle der skal bruge
   det. Bygget-og-korrekt-men-utilgængeligt er husets hyppigste fejlklasse
   (`read_attachment` 6/9, billedværktøjerne 13/9, visningsværktøjerne 27/10).
3. **Virkelig** — filen kan læses tilbage med sit eget format-bibliotek.
   En fil med en størrelse er ikke en fil med et indhold; den første version
   skrev overskrifter og nul rækker uden at fejle.
"""
from __future__ import annotations

import pytest

from core.identity.workspace_context import user_context
from core.services.permission_engine import allowed_tools
from core.tools.tool_scoping import allowed_tool_names

BJORN = "1246415163603816499"
MICHELLE = "1313522677369143429"

OPRET = {"pdf_create", "docx_create", "xlsx_create", "pptx_create"}
ALLE = OPRET | {"pdf_read"}


# ── 1. findes ─────────────────────────────────────────────────────────────

def test_alle_fem_staar_i_kataloget():
    """Schemaet er det eneste der gør et værktøj synligt for modellen."""
    from core.tools.simple_tools_definitions import TOOL_DEFINITIONS

    navne = {
        ((d.get("function") or {}).get("name") or d.get("name") or "")
        for d in TOOL_DEFINITIONS
    }
    mangler = ALLE - navne
    assert not mangler, f"i kataloget mangler: {sorted(mangler)}"


def test_hver_har_en_handler():
    """Et schema uden handler svarer «unknown tool» — død men synlig."""
    from core.tools.simple_tools import _TOOL_HANDLERS

    mangler = {n for n in ALLE if n not in _TOOL_HANDLERS}
    assert not mangler, f"uden handler: {sorted(mangler)}"


# ── 2. kan nås ────────────────────────────────────────────────────────────

@pytest.mark.parametrize("scope", ["chat", "code"])
def test_ejeren_kan_oprette_i_begge_modes(scope):
    with user_context(discord_id=BJORN):
        ud = allowed_tool_names(role="owner", scope=scope, all_names=ALLE)
    assert OPRET <= ud, f"owner/{scope} mistede: {sorted(OPRET - ud)}"


@pytest.mark.parametrize("scope", ["chat", "code"])
def test_husstanden_kan_oprette_i_begge_modes(scope):
    """Michelle skal kunne lave et regneark — ikke kun læse Bjørns."""
    with user_context(discord_id=MICHELLE):
        ud = allowed_tool_names(role="partner", scope=scope, all_names=ALLE)
    assert OPRET <= ud, f"partner/{scope} mistede: {sorted(OPRET - ud)}"


@pytest.mark.parametrize("scope", ["chat", "code"])
def test_pdf_read_er_kun_for_ejeren(scope):
    """`pdf_read` læser en VILKÅRLIG sti — samme fil-adgang som `read_file`."""
    with user_context(discord_id=BJORN):
        ejer = allowed_tool_names(role="owner", scope=scope, all_names=ALLE)
    with user_context(discord_id=MICHELLE):
        husstand = allowed_tool_names(role="partner", scope=scope, all_names=ALLE)
    assert "pdf_read" in ejer
    assert "pdf_read" not in husstand, "en member må ikke kunne læse vilkårlige stier"


def test_permission_engine_giver_husstanden_opret():
    """Den ANDEN kilde: for member/partner er permission_engine loft OG gulv.

    Målt 27/10-2026: scope-listen alene gav partner nul værktøjer, fordi
    `allowed_tools(role, mode)` returnerer sin egen liste. Begge steder skal
    rettes — ellers ser det rigtigt ud i én måling og virker ikke i praksis.
    """
    for mode in ("chat", "code"):
        ud = allowed_tools(role="partner", mode=mode)
        assert OPRET <= set(ud), f"permission_engine/{mode} mistede: {sorted(OPRET - set(ud))}"
        assert "pdf_read" not in set(ud), "pdf_read må ikke åbne for husstanden"


# ── 3. virkelig ───────────────────────────────────────────────────────────

@pytest.fixture()
def som_bjorn():
    with user_context(discord_id=BJORN):
        yield


def _sti(navn: str):
    from core.runtime.workspace_paths import published_file_path
    return published_file_path(navn, opret_mappe=False)


def test_pdf_kan_laeses_tilbage(som_bjorn):
    from core.tools.document_create_tools import _exec_pdf_create

    r = _exec_pdf_create({
        "titel": "Rapport",
        "indhold": "# Konklusion\nIngen drift fundet.",
        "filnavn": "test-doc-pdf",
    })
    assert r["status"] == "ok", r

    from pypdf import PdfReader
    tekst = (PdfReader(str(_sti("test-doc-pdf.pdf"))).pages[0].extract_text() or "")
    assert "Konklusion" in tekst
    assert "Ingen drift fundet" in tekst


def test_word_kan_laeses_tilbage(som_bjorn):
    from core.tools.document_create_tools import _exec_docx_create

    r = _exec_docx_create({
        "titel": "Notat",
        "indhold": "# Afsnit\nBrødtekst her.",
        "filnavn": "test-doc-docx",
    })
    assert r["status"] == "ok", r

    import docx
    tekst = " ".join(p.text for p in docx.Document(str(_sti("test-doc-docx.docx"))).paragraphs)
    assert "Notat" in tekst
    assert "Brødtekst her" in tekst


def test_praesentation_kan_laeses_tilbage(som_bjorn):
    from core.tools.document_create_tools import _exec_pptx_create

    r = _exec_pptx_create({
        "titel": "Præsentation",
        "slides": [{"titel": "Slide et", "punkter": ["punkt"]}],
        "filnavn": "test-doc-pptx",
    })
    assert r["status"] == "ok", r

    from pptx import Presentation
    prs = Presentation(str(_sti("test-doc-pptx.pptx")))
    titler = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert "Præsentation" in titler
    assert "Slide et" in titler


@pytest.mark.parametrize("noegle", ["raekker", "rækker", "rows"])
def test_regneark_faar_data_uanset_stavemaade(som_bjorn, noegle):
    """æ/ø/å i JSON er en tavs fælde.

    Målt 27/10-2026: schemaet skriver `rækker`, den kaldende skriver `raekker`,
    og den første version skrev en fil med OVERSKRIFTER OG NUL RÆKKER — ingen
    fejl, bare en tom tabel der ser ud som en tabel. Alle tre stavemåder skal
    give samme indhold.
    """
    from openpyxl import load_workbook

    from core.tools.document_create_tools import _exec_xlsx_create

    navn = f"test-doc-xlsx-{noegle}"
    r = _exec_xlsx_create({
        "titel": "Ark",
        "ark": [{"navn": "Data", "kolonner": ["Navn", "Værdi"],
                 noegle: [["a", 1], ["b", 2]]}],
        "filnavn": navn,
    })
    assert r["status"] == "ok", r

    ws = load_workbook(str(_sti(f"{navn}.xlsx"))).active
    raekker = list(ws.iter_rows(values_only=True))
    assert raekker[0] == ("Navn", "Værdi"), f"header mangler: {raekker[:1]}"
    assert len(raekker) == 3, f"data-rækkerne forsvandt: {raekker}"
    assert raekker[1] == ("a", 1)
    assert raekker[2] == ("b", 2)


def test_filnavn_faar_endelsen_tilfoejet(som_bjorn):
    """En fil uden endelse kan ikke åbnes af det program den hører til."""
    from core.tools.document_create_tools import _exec_pdf_create

    r = _exec_pdf_create({"titel": "Uden endelse", "indhold": "tekst",
                          "filnavn": "test-doc-uden-endelse"})
    assert r["status"] == "ok", r
    assert r["filename"].endswith(".pdf")
